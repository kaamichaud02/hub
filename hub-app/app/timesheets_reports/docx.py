"""Génération du rapport hebdomadaire au format Word (.docx). Port de
suivi_temps/timesheets/reports/docx.py — logique python-docx inchangée,
seule la source des entrées et le type de réponse HTTP changent."""
from io import BytesIO

from fastapi import Response
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX

from .. import timesheets_services as services


def _styled_para(doc, text=None, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if text:
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.bold = bold
    return p


def _add_run(para, text, bold=False, highlight=False):
    run = para.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.bold = bold
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def render_docx_report(data: dict, monday: str) -> Response:
    """Construit la réponse HTTP contenant le rapport Word de la semaine.

    `data` provient de `timesheets_reports.common.build_week_report_data`.
    """
    user = data["user"]
    monday_date = data["monday_date"]
    sunday_date = data["sunday_date"]
    pairs = data["timesheets"]

    doc = Document()

    section = doc.sections[0]
    header_para = section.header.paragraphs[0]
    header_para.text = f"Rapport des feuilles de temps de {user.first_name} {user.last_name}"
    header_run = header_para.runs[0]
    header_run.font.name = 'Arial'
    header_run.font.size = Pt(12)
    header_run.bold = True

    heading = doc.add_heading(
        f"Semaine du {monday_date.strftime('%d/%m/%Y')} au {sunday_date.strftime('%d/%m/%Y')}",
        level=2,
    )
    heading.paragraph_format.space_after = Pt(0)

    _styled_para(doc)
    _styled_para(doc, f"Total: {data['total_hours']} h {data['total_minutes']} min")
    _styled_para(doc)

    for ts, entries in pairs:
        hours, minutes = services.timesheet_duration(ts.date, entries)

        p = _styled_para(doc)
        _add_run(p, f"Date: {ts.date.strftime('%d/%m/%Y')} (Total: ", bold=True)
        _add_run(p, f"{hours} h {minutes} min", bold=True, highlight=True)
        _add_run(p, ")", bold=True)

        if entries:
            for entry in entries:
                _styled_para(doc, f"{entry.start_time.strftime('%H:%M')} - {entry.end_time.strftime('%H:%M')}")
        else:
            _styled_para(doc, "Aucune entrée horaire")

        _styled_para(doc)
        _styled_para(doc, "Résumé:")

        if ts.summary:
            _styled_para(doc)
            for line in ts.summary.split('\n'):
                if line.strip():
                    _styled_para(doc, line.strip())
        else:
            _styled_para(doc)

        _styled_para(doc)
        _styled_para(doc)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="rapport_semaine_{monday}.docx"'},
    )
